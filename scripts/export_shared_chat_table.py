#!/usr/bin/env python3
"""Export a Markdown table from one explicitly selected Open WebUI share."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def split_row(line: str) -> list[str]:
    text = line.strip().strip("|")
    return [cell.strip().replace(r"\|", "|") for cell in re.split(r"(?<!\\)\|", text)]


def extract_table(content: str) -> list[list[str]]:
    lines = content.splitlines()
    for index in range(len(lines) - 2):
        if not lines[index].lstrip().startswith("|"):
            continue
        separator = split_row(lines[index + 1])
        if not separator or not all(re.fullmatch(r":?-{3,}:?", item) for item in separator):
            continue
        header = split_row(lines[index])
        rows = [header]
        for line in lines[index + 2:]:
            if not line.lstrip().startswith("|"):
                break
            row = split_row(line)
            if len(row) != len(header):
                raise ValueError("Markdown table has inconsistent column counts")
            rows.append(row)
        if len(rows) > 1:
            return rows
    raise ValueError("No Markdown table found in shared answer")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("share_id")
    parser.add_argument("output")
    parser.add_argument("--db", default="/app/backend/data/webui.db")
    args = parser.parse_args()
    with sqlite3.connect(f"file:{args.db}?mode=ro", uri=True) as db:
        row = db.execute("SELECT chat FROM shared_chat WHERE id=?", (args.share_id,)).fetchone()
    if not row:
        raise SystemExit("Shared chat was not found")
    chat = json.loads(row[0])
    messages = chat.get("history", {}).get("messages", {})
    assistant_answers = [message.get("content", "") for message in messages.values() if message.get("role") == "assistant"]
    if not assistant_answers:
        raise SystemExit("Shared chat has no assistant answer")
    rows = extract_table(assistant_answers[-1])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Câu hỏi tương tự"
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
    for index, width in enumerate((9, 82, 58, 50)[:len(rows[0])], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = Path(args.output)
    if not output.parent.is_dir():
        raise SystemExit("Output directory does not exist")
    workbook.save(output)
    print(f"Exported {len(rows) - 1} data rows")


if __name__ == "__main__":
    main()
