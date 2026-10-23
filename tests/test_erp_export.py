"""Check downloadable ERP files, including spreadsheet text safety."""

import importlib.util
from io import BytesIO
import json
from pathlib import Path
import shutil
import unittest

from erp_ai.export_files import blocks_from_rows, build_downloads, requested_formats
from erp_ai.rich_ui import render_erp_embed


BLOCKS = [{
    "type": "table", "title": "Nhập kho", "columns": [
        {"key": "so_ct", "label": "Số chứng từ"},
        {"key": "sl_nhap", "label": "Số lượng"},
    ],
    "rows": [{"so_ct": "=1+1", "sl_nhap": 2}],
    "total_rows": 9, "has_more": True,
}]


class ExportTests(unittest.TestCase):
    def test_requested_formats(self):
        self.assertEqual(requested_formats("xem báo cáo"), ("xlsx", "docx", "pdf"))
        self.assertEqual(requested_formats("xuất Word và PDF"), ("xlsx", "docx", "pdf"))

    def test_blocks_from_verified_rows(self):
        blocks = blocks_from_rows({"rows": [{"ma_vt": "A", "sl": 2}], "truncated": True})
        self.assertTrue(blocks[0]["has_more"])
        self.assertEqual(blocks[0]["columns"][0]["key"], "ma_vt")
        self.assertEqual(blocks_from_rows({"rows": ["unstructured"]}), [])
        self.assertEqual(blocks_from_rows({"status": "invalid_result", "rows": [{"dvt": "PCS"}]}), [])

    def test_blocks_from_rows_cleans_trailing_zeros(self):
        payload = {
            "status": "success",
            "summary": {
                "tong_so_luong": "229140.00000000000000000000",
                "tong_doanh_thu": "457656389.82400000000000000000",
            },
            "rows": [{
                "ma_vt": "QC7-5146-000",
                "so_luong": "229140.00000000000000000000",
                "doanh_thu": "175446405.12000000000000000000",
                "total": 1,
            }],
        }
        blocks = blocks_from_rows(payload)
        self.assertEqual(blocks[0]["text"], "Thống kê ERP — Tổng số lượng: 229140, Tổng doanh thu: 457656389.824")
        self.assertEqual(blocks[1]["rows"][0]["so_luong"], "229140")
        self.assertEqual(blocks[1]["rows"][0]["doanh_thu"], "175446405.12")
        self.assertNotIn("total", [c["key"] for c in blocks[1]["columns"]])

    def test_existing_erp_report_shows_summary_and_receipt_chart(self):
        payload = {
            "status": "success", "report": {"title": "Bảng kê phiếu nhập kho"},
            "summary": {"total_records": 3, "total_quantity": 9},
            "rows": [
                {"ngay_ct": "2026-10-20", "sl_nhap": 2, "dvt": "cái"},
                {"ngay_ct": "2026-10-21", "sl_nhap": 3, "dvt": "cái"},
            ], "total_rows": 3,
        }
        blocks = blocks_from_rows(payload)
        self.assertEqual([block["type"] for block in blocks], ["summary", "table", "chart"])
        self.assertIn("Tổng số lượng: 9", blocks[0]["text"])
        self.assertEqual(blocks[1]["title"], "Bảng kê phiếu nhập kho")
        self.assertEqual(blocks[2]["scope"], "2 dòng đang hiển thị")
        self.assertEqual(blocks[2]["points"][1]["quantity"], 3)
        html = render_erp_embed({"blocks": blocks})
        self.assertIn("Số lượng nhập theo ngày", html)
        self.assertIn("Tổng số lượng: 9", html)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl") and importlib.util.find_spec("docx"), "Office libraries unavailable")
    def test_office_files_and_embed(self):
        from openpyxl import load_workbook
        from docx import Document

        downloads = build_downloads(BLOCKS, ("xlsx", "docx"))
        self.assertEqual([item["name"] for item in downloads], ["bao-cao-erp.xlsx", "bao-cao-erp.docx"])
        sheet = load_workbook(BytesIO(downloads[0]["data"])).active
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertIn("Chỉ xuất 1 dòng", sheet["A4"].value)
        document = Document(BytesIO(downloads[1]["data"]))
        self.assertEqual(document.tables[0].cell(1, 0).text, "=1+1")
        html = render_erp_embed({"blocks": BLOCKS}, downloads)
        self.assertIn("Tải XLSX", html)
        self.assertIn("Tải DOCX", html)
        self.assertIn("data-download-index", html)

    @unittest.skipUnless(importlib.util.find_spec("docx") and shutil.which("libreoffice"), "PDF converter unavailable")
    def test_pdf_file(self):
        downloads = build_downloads(BLOCKS, ("pdf",))
        self.assertEqual(downloads[0]["name"], "bao-cao-erp.pdf")
        self.assertTrue(downloads[0]["data"].startswith(b"%PDF-"))

    @unittest.skipUnless(importlib.util.find_spec("fastapi") and importlib.util.find_spec("openpyxl"), "Open WebUI libraries unavailable")
    def test_existing_erp_result_remains_visible_to_model(self):
        staged = Path("/tmp/htmp-erp-export-deploy/htmp_postgres_query.py")
        source = staged if staged.is_file() else Path(__file__).resolve().parents[1] / "functions" / "htmp_postgres_query.py"
        spec = importlib.util.spec_from_file_location("erp_export_tool_test", source)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        tool = module.Tools()
        tool._ask_erp_data = lambda *args: json.dumps({"rows": [{"so_ct": "PN-1", "sl_nhap": 2}], "row_count": 1})
        response, context = tool.ask_erp("liệt kê nhập kho")
        self.assertEqual(context["rows"][0]["so_ct"], "PN-1")
        self.assertEqual(context["download_formats"], ["xlsx", "docx", "pdf"])
        self.assertIn("Tải XLSX", response.body.decode())


if __name__ == "__main__":
    unittest.main()
